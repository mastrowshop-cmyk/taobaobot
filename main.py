from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from openpyxl import load_workbook
import json

from database import Base, engine, SessionLocal
from models import User, Parcel, Settings
from telegram_auth import verify_init_data, extract_tg_user
from telegram_notify import tg_send
from config import ALLOWED_ORIGINS

Base.metadata.create_all(bind=engine)

app = FastAPI(title="Nataobao API")

origins = ["*"] if ALLOWED_ORIGINS == "*" else [x.strip() for x in ALLOWED_ORIGINS.split(",") if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def db() -> Session:
    return SessionLocal()

def get_or_init_settings(session: Session) -> Settings:
    s = session.query(Settings).first()
    if not s:
        s = Settings()
        session.add(s)
        session.commit()
        session.refresh(s)
    return s

def auth_user(init_data: str) -> User:
    ok, _ = verify_init_data(init_data)
    if not ok:
        raise HTTPException(status_code=401, detail="Bad initData")

    tg_user = extract_tg_user(init_data)
    if not tg_user.get("id"):
        raise HTTPException(status_code=401, detail="No user in initData")

    tg_id = int(tg_user["id"])
    session = db()
    u = session.query(User).filter(User.tg_id == tg_id).first()
    if not u:
        raise HTTPException(status_code=403, detail="Not registered")
    return u

@app.get("/health")
def health():
    return {"ok": True}

@app.get("/")
def root():
    # не обязательно, но пусть будет красиво
    return {"ok": True, "service": "nataobao-api"}

# --------- WEBAPP API ---------

@app.post("/me")
def me(payload: dict):
    init_data = payload.get("initData", "")
    u = auth_user(init_data)

    session = db()
    s = get_or_init_settings(session)

    return {
        "tg_id": u.tg_id,
        "role": u.role,
        "code": u.code,
        "first_name": u.first_name,
        "last_name": u.last_name,
        "support_username": s.support_username
    }

@app.post("/profile/update")
def profile_update(payload: dict):
    init_data = payload.get("initData", "")
    first_name = (payload.get("first_name") or "").strip()
    last_name = (payload.get("last_name") or "").strip()

    if not first_name or not last_name:
        raise HTTPException(status_code=400, detail="Empty name")

    u = auth_user(init_data)
    session = db()
    u2 = session.query(User).filter(User.tg_id == u.tg_id).first()
    u2.first_name = first_name
    u2.last_name = last_name
    session.commit()
    return {"ok": True}

@app.post("/parcels/list")
def parcels_list(payload: dict):
    init_data = payload.get("initData", "")
    u = auth_user(init_data)

    session = db()
    parcels = session.query(Parcel).filter(Parcel.user_code == u.code).order_by(Parcel.id.desc()).all()
    return [{
        "id": p.id,
        "title": p.title,
        "track": p.track,
        "status": p.status,
        "weight_kg": p.weight_kg,
        "price_rub": p.price_rub,
        "paid": p.paid
    } for p in parcels]

@app.post("/payment/info")
def payment_info(payload: dict):
    init_data = payload.get("initData", "")
    _ = auth_user(init_data)

    session = db()
    s = get_or_init_settings(session)
    return {"pay_text": s.pay_text}

# --------- OPERATOR (Якутск) ---------

@app.post("/operator/weigh")
def operator_weigh(payload: dict):
    init_data = payload.get("initData", "")
    operator = auth_user(init_data)
    if operator.role not in ("operator", "admin"):
        raise HTTPException(status_code=403, detail="Forbidden")

    code = (payload.get("code") or "").strip()
    parcel_id = int(payload.get("parcel_id") or 0)
    weight_kg = float(payload.get("weight_kg") or 0)

    if not code or parcel_id <= 0 or weight_kg <= 0:
        raise HTTPException(status_code=400, detail="Bad input")

    session = db()
    s = get_or_init_settings(session)

    p = session.query(Parcel).filter(Parcel.id == parcel_id, Parcel.user_code == code).first()
    if not p:
        raise HTTPException(status_code=404, detail="Parcel not found")

    price = weight_kg * float(s.price_per_kg)

    p.weight_kg = weight_kg
    p.price_rub = price
    p.status = "pay"
    session.commit()

    client = session.query(User).filter(User.code == code).first()
    if client:
        tg_send(client.tg_id, f"📦 Nataobao: посылка взвешена.\nВес: {weight_kg} кг\nК оплате: {round(price)} ₽")

    return {"ok": True, "price_rub": price}

# --------- ADMIN ---------

def require_admin(u: User):
    if u.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")

@app.post("/admin/settings/get")
def admin_settings_get(payload: dict):
    init_data = payload.get("initData", "")
    u = auth_user(init_data)
    require_admin(u)

    session = db()
    s = get_or_init_settings(session)
    return {
        "price_per_kg": s.price_per_kg,
        "pay_text": s.pay_text,
        "support_username": s.support_username
    }

@app.post("/admin/settings/save")
def admin_settings_save(payload: dict):
    init_data = payload.get("initData", "")
    u = auth_user(init_data)
    require_admin(u)

    session = db()
    s = get_or_init_settings(session)

    if "price_per_kg" in payload:
        s.price_per_kg = float(payload["price_per_kg"])
    if "pay_text" in payload:
        s.pay_text = payload["pay_text"]
    if "support_username" in payload:
        s.support_username = (payload["support_username"] or "").replace("@", "").strip() or s.support_username

    session.commit()
    return {"ok": True}

@app.post("/admin/users/pending")
def admin_users_pending(payload: dict):
    init_data = payload.get("initData", "")
    u = auth_user(init_data)
    require_admin(u)

    session = db()
    users = session.query(User).filter(User.role == "pending").order_by(User.id.desc()).all()
    return [{
        "tg_id": x.tg_id,
        "code": x.code,
        "first_name": x.first_name,
        "last_name": x.last_name,
        "username": x.username
    } for x in users]

@app.post("/admin/users/approve")
def admin_users_approve(payload: dict):
    init_data = payload.get("initData", "")
    u = auth_user(init_data)
    require_admin(u)

    tg_id = int(payload.get("tg_id") or 0)
    if tg_id <= 0:
        raise HTTPException(status_code=400, detail="Bad tg_id")

    session = db()
    usr = session.query(User).filter(User.tg_id == tg_id).first()
    if not usr:
        raise HTTPException(status_code=404, detail="Not found")

    usr.role = "client"
    session.commit()

    tg_send(tg_id, "✅ Nataobao: ваша заявка одобрена! Откройте приложение в боте.")
    return {"ok": True}

@app.post("/admin/users/reject")
def admin_users_reject(payload: dict):
    init_data = payload.get("initData", "")
    u = auth_user(init_data)
    require_admin(u)

    tg_id = int(payload.get("tg_id") or 0)
    if tg_id <= 0:
        raise HTTPException(status_code=400, detail="Bad tg_id")

    session = db()
    usr = session.query(User).filter(User.tg_id == tg_id).first()
    if not usr:
        raise HTTPException(status_code=404, detail="Not found")

    session.delete(usr)
    session.commit()

    tg_send(tg_id, "❌ Nataobao: заявка отклонена.")
    return {"ok": True}

@app.post("/admin/excel/import")
async def admin_excel_import(initData: str = Form(""), file: UploadFile = File(...)):
    # initData в form-field, файл — multipart
    u = auth_user(initData)
    require_admin(u)

    content = await file.read()
    with open("import.xlsx", "wb") as f:
        f.write(content)

    wb = load_workbook("import.xlsx")
    ws = wb.active

    session = db()
    notified_codes = set()

    # Excel формат (со 2-й строки):
    # CODE | TITLE | TRACK(optional) | STATUS(optional)
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = (row[0] or "").strip() if row[0] else ""
        title = (row[1] or "").strip() if row[1] else ""
        track = str(row[2]).strip() if row[2] else None
        status = str(row[3]).strip() if row[3] else "yakutsk"

        if not code or not title:
            continue

        session.add(Parcel(
            user_code=code,
            title=title,
            track=track,
            status=status
        ))
        notified_codes.add(code)

    session.commit()

    # уведомим клиентов
    for code in notified_codes:
        client = session.query(User).filter(User.code == code).first()
        if client:
            tg_send(client.tg_id, "📦 Nataobao: добавлена новая доставка. Проверьте «Мои доставки».")

    return {"ok": True}
